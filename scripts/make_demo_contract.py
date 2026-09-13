"""Generate a small demo service agreement DOCX for screenshots / demos."""
from docx import Document
from docx.shared import Pt

d = Document()
style = d.styles["Normal"]
style.font.name = "Georgia"
style.font.size = Pt(11)

d.add_heading("MASTER SERVICES AGREEMENT", level=1)
d.add_paragraph(
    "This Master Services Agreement (the \"Agreement\") is entered into as of "
    "March 4, 2026 (the \"Effective Date\"), by and between Northwind Analytics LLC, "
    "a Delaware limited liability company with its principal place of business in "
    "Austin, Texas (\"Provider\"), and Bluepeak Retail Inc., a California corporation "
    "(\"Client\"). Provider and Client are each a \"Party\" and together the \"Parties\"."
)
d.add_heading("1. Services", level=2)
d.add_paragraph(
    "Provider shall supply the data analytics and dashboard services described in one or "
    "more mutually executed statements of work (each, an \"SOW\"). Each SOW shall be "
    "governed by this Agreement."
)
d.add_heading("2. License Grant", level=2)
d.add_paragraph(
    "Subject to Client's payment obligations, Provider grants Client a limited, "
    "non-exclusive, non-transferable, non-sublicensable license during the Term to "
    "access and use the Provider software solely for Client's internal business purposes."
)
d.add_heading("3. Term and Renewal", level=2)
d.add_paragraph(
    "The initial term of this Agreement is twenty-four (24) months beginning on the "
    "Effective Date (the \"Initial Term\"). This Agreement shall automatically renew for "
    "successive twelve (12) month renewal terms unless either Party provides written "
    "notice of non-renewal at least sixty (60) days prior to the end of the then-current term."
)
d.add_heading("4. Termination", level=2)
d.add_paragraph(
    "Either Party may terminate this Agreement for convenience upon ninety (90) days' "
    "prior written notice to the other Party. Either Party may terminate this Agreement "
    "immediately upon written notice if the other Party materially breaches this "
    "Agreement and fails to cure such breach within thirty (30) days after receipt of "
    "written notice thereof."
)
d.add_heading("5. Non-Compete", level=2)
d.add_paragraph(
    "During the Term and for a period of twelve (12) months thereafter, Client shall not "
    "directly or indirectly engage, contract with, or solicit the services of any "
    "competitor of Provider for the purpose of receiving analytics services that are "
    "materially similar to the Services."
)
d.add_heading("6. Exclusivity", level=2)
d.add_paragraph(
    "During the Term, Client agrees that Provider shall be the exclusive provider of "
    "managed data analytics services to Client, and Client shall not retain any third "
    "party to provide services that are the same as or substantially similar to the Services."
)
d.add_heading("7. Assignment", level=2)
d.add_paragraph(
    "Neither Party may assign this Agreement or any of its rights or obligations "
    "hereunder without the prior written consent of the other Party, except that either "
    "Party may assign this Agreement without consent in connection with a merger or sale "
    "of substantially all of its assets."
)
d.add_heading("8. Audit Rights", level=2)
d.add_paragraph(
    "Client shall have the right, upon thirty (30) days' prior written notice and no more "
    "than once per calendar year, to audit Provider's records relevant to the charges "
    "under this Agreement during normal business hours."
)
d.add_heading("9. Limitation of Liability", level=2)
d.add_paragraph(
    "NOTWITHSTANDING ANYTHING TO THE CONTRARY, IN NO EVENT SHALL EITHER PARTY BE LIABLE "
    "FOR ANY INDIRECT, INCIDENTAL, SPECIAL, CONSEQUENTIAL OR PUNITIVE DAMAGES. EACH "
    "PARTY'S AGGREGATE LIABILITY ARISING OUT OF OR RELATED TO THIS AGREEMENT SHALL NOT "
    "EXCEED THE TOTAL FEES PAID BY CLIENT TO PROVIDER IN THE TWELVE (12) MONTHS "
    "PRECEDING THE CLAIM."
)
d.add_heading("10. Insurance", level=2)
d.add_paragraph(
    "Provider shall maintain, at its own expense, commercial general liability insurance "
    "with limits of not less than $2,000,000 per occurrence and professional liability "
    "(errors and omissions) insurance with limits of not less than $5,000,000 in the aggregate."
)
d.add_heading("11. Governing Law", level=2)
d.add_paragraph(
    "This Agreement and any dispute arising out of or in connection with it shall be "
    "governed by and construed in accordance with the laws of the State of Delaware, "
    "without regard to its conflict of laws principles. The Parties consent to the "
    "exclusive jurisdiction of the state and federal courts located in New Castle County, Delaware."
)

d.save(r"D:\ContractIQ\storage\demo_services_agreement.docx")
print("saved")
