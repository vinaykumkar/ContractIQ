"""Document text extraction for PDF, DOCX and TXT.

Safety model: documents are untrusted DATA. We only extract text with local
libraries (PyMuPDF, python-docx); no content is ever executed, no embedded
objects/macros are opened, and OCR is explicitly out of scope (a text-less PDF
is reported as `ocr_required`, never silently faked).

Every parser returns a ParsedDocument; all failures raise the domain
exceptions from app.core.exceptions.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from ..core.exceptions import (
    CorruptDocument,
    EmptyDocument,
    EncryptedDocument,
    OcrRequiredError,
    ParseError,
    UnsupportedFileType,
)

# below this many extractable characters a PDF is treated as text-less
PDF_MIN_TEXT_CHARS = 20


@dataclass
class ParsedDocument:
    """Standard parser output for the whole pipeline."""

    filename: str
    original_filename: str
    file_type: str  # "pdf" | "docx" | "txt"
    text: str
    character_count: int
    page_count: int | None = None
    metadata: dict = field(default_factory=dict)
    warnings: list[str] = field(default_factory=list)
    ocr_required: bool = False

    @property
    def is_empty(self) -> bool:
        return not self.text.strip()


def parse_document(path: Path, original_filename: str | None = None) -> ParsedDocument:
    """Dispatch on the (already validated) file extension."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(path, original_filename or path.name)
    if ext == ".docx":
        return parse_docx(path, original_filename or path.name)
    if ext == ".txt":
        return parse_txt(path, original_filename or path.name)
    raise UnsupportedFileType(f"Cannot parse '{ext or '(no extension)'}' documents.")


# ------------------------------------------------------------------------ PDF

def parse_pdf(path: Path, original_filename: str) -> ParsedDocument:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover
        raise ParseError("PyMuPDF is not installed.") from exc

    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise CorruptDocument(f"PDF could not be opened: {type(exc).__name__}") from exc

    try:
        try:
            if doc.needs_pass:
                raise EncryptedDocument(
                    "PDF is password-protected; encrypted documents are not processed."
                )
        except EncryptedDocument:
            raise
        except Exception as exc:
            raise CorruptDocument(f"PDF structure could not be read: {type(exc).__name__}") from exc

        page_count = int(doc.page_count)
        if page_count == 0:
            raise EmptyDocument("PDF contains no pages.")
        chunks: list[str] = []
        for page in doc:
            try:
                chunks.append(page.get_text("text"))
            except Exception as exc:  # malformed page content
                raise CorruptDocument(f"PDF page could not be extracted: {type(exc).__name__}") from exc
        text = "\n".join(chunks)
    finally:
        doc.close()

    warnings: list[str] = []
    if not text.strip() or len(text.strip()) < PDF_MIN_TEXT_CHARS:
        # pages exist but carry no text layer -> scanned/image-only; OCR stays optional
        warnings.append(
            "No usable text layer found - the PDF is likely scanned/image-only. "
            "OCR is optional and not enabled."
        )
        raise OcrRequiredError(warnings[0])

    return ParsedDocument(
        filename=path.name,
        original_filename=original_filename,
        file_type="pdf",
        text=text,
        character_count=len(text),
        page_count=page_count,
        metadata={"parser": "pymupdf", "page_count": page_count},
        warnings=warnings,
        ocr_required=False,
    )


# ----------------------------------------------------------------------- DOCX

def parse_docx(path: Path, original_filename: str) -> ParsedDocument:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise ParseError("python-docx is not installed.") from exc

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise CorruptDocument(f"DOCX could not be opened: {type(exc).__name__}") from exc

    parts: list[str] = []
    for para in d.paragraphs:
        parts.append(para.text)
    for table in d.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise EmptyDocument("DOCX contains no readable text.")

    return ParsedDocument(
        filename=path.name,
        original_filename=original_filename,
        file_type="docx",
        text=text,
        character_count=len(text),
        page_count=None,  # DOCX has no fixed pagination
        metadata={"parser": "python-docx", "tables": len(d.tables), "paragraphs": len(d.paragraphs)},
        warnings=[],
        ocr_required=False,
    )


# ------------------------------------------------------------------------ TXT

def parse_txt(path: Path, original_filename: str) -> ParsedDocument:
    raw = path.read_bytes()
    if not raw.strip():
        raise EmptyDocument("TXT file is empty.")

    text = None
    warnings: list[str] = []
    encoding_used = "utf-8"
    for enc in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            encoding_used = enc
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise CorruptDocument("TXT file could not be decoded with any supported encoding.")
    if encoding_used != "utf-8":
        warnings.append(f"Decoded with fallback encoding '{encoding_used}'.")

    return ParsedDocument(
        filename=path.name,
        original_filename=original_filename,
        file_type="txt",
        text=text,
        character_count=len(text),
        page_count=None,
        metadata={"parser": "builtin", "encoding": encoding_used},
        warnings=warnings,
        ocr_required=False,
    )
